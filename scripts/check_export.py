"""Verify a generated proposal workbook: no blank sums, and the arithmetic adds up.

    python scripts/check_export.py                # newest estimate in the database
    python scripts/check_export.py 3              # a specific estimate
    python scripts/check_export.py --out /tmp/kp.xlsx

Exits non-zero and prints every failure, so it can gate a deploy.

Why this exists: openpyxl writes formulas with an *empty* cached result, so a
workbook whose formulas are all correct still reads as entirely blank to
anything that does not recalculate — LibreOffice, Google Sheets, a thumbnail
preview, and `load_workbook(data_only=True)`. The export now fills that cache
in, and this script is what proves it, on real estimate data rather than a
fixture.
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "backend"))

import openpyxl  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from app.api.routes import _layout  # noqa: E402
from app.db import SessionLocal  # noqa: E402
from app.models import Estimate, Project  # noqa: E402
from app.services.estimate.store import draft_from_estimate  # noqa: E402
from app.services.export.xlsx import (  # noqa: E402
    LAST_COL,
    MONEY,
    MONEY_STRONG,
    QTY,
    ExportMeta,
    export_estimate,
)
from app.services.rules.totals import compute_totals  # noqa: E402
from app.services.validation.validators import validate  # noqa: E402

TOTAL_PREFIXES = ("Разом", "Загальний рахунок", "Аванс", "Залишок", "Безготівковий")
CENT = 0.011  # a hundredth, plus float slack


class Checker:
    def __init__(self) -> None:
        self.failures: list[str] = []
        self.checks = 0

    def ok(self, condition: bool, message: str) -> bool:
        self.checks += 1
        if not condition:
            self.failures.append(message)
        return condition

    def report(self) -> int:
        print(f"\n{self.checks} checks run.")
        if not self.failures:
            print("PASS — every sum is present and the arithmetic agrees.")
            return 0
        print(f"FAIL — {len(self.failures)} problem(s):")
        for f in self.failures:
            print("  •", f)
        return 1


def build_workbook(estimate_id: int | None, out: Path) -> tuple[Path, str]:
    session = SessionLocal()
    try:
        query = session.query(Estimate).order_by(Estimate.id.desc())
        estimate = session.get(Estimate, estimate_id) if estimate_id else query.first()
        if estimate is None:
            raise SystemExit(
                "No estimate found. Pass an id, or build one through the API first."
            )
        project = session.get(Project, estimate.project_id)
        layout = _layout()
        draft = draft_from_estimate(estimate)
        totals = compute_totals(draft, layout.titles(), estimate.settings)
        report = validate(draft, totals)
        export_estimate(
            out,
            draft,
            totals,
            meta=ExportMeta(
                client_name=project.client_name if project else "",
                address=project.address if project else "",
                project_name=project.name if project else "",
                date=dt.date.today().strftime("%d.%m.%Y"),
            ),
            section_titles=layout.titles(),
            section_order=layout.order,
            report=report,
        )
        label = f"estimate {estimate.id} ({project.name if project else '—'})"
        return out, label
    finally:
        session.close()


def label_of(ws, row: int) -> str:
    """The text that names a row, whether it sits in column A or column B."""
    for col in (1, 2):
        value = ws.cell(row, col).value
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def check(path: Path, c: Checker) -> None:
    live = openpyxl.load_workbook(path)          # formulas
    cached = openpyxl.load_workbook(path, data_only=True)  # what a viewer sees
    ws, vs = live["Кошторис"], cached["Кошторис"]

    line_rows: list[int] = []
    total_rows: list[int] = []
    header_rows: list[int] = []

    for row in range(1, ws.max_row + 1):
        first = ws.cell(row, 1).value
        if first == "#":
            header_rows.append(row)
        elif isinstance(first, int) and ws.cell(row, 2).value:
            line_rows.append(row)
        elif str(label_of(ws, row)).startswith(TOTAL_PREFIXES):
            total_rows.append(row)

    c.ok(bool(line_rows), "no item rows found in the proposal sheet")
    c.ok(bool(total_rows), "no «Разом …» rows found in the proposal sheet")

    # --- 1. nothing in a Сума column is blank --------------------------------
    for row in line_rows + total_rows:
        name = label_of(ws, row)[:40]
        formula = ws.cell(row, LAST_COL).value
        value = vs.cell(row, LAST_COL).value
        c.ok(
            isinstance(formula, str) and formula.startswith("="),
            f"F{row} ({name}): no formula, got {formula!r}",
        )
        c.ok(
            value is not None and value != "",
            f"F{row} ({name}): sum is blank when the file is read without recalculating",
        )

    # --- 2. the arithmetic ---------------------------------------------------
    for row in line_rows:
        qty, price = vs.cell(row, 3).value, vs.cell(row, 5).value
        total = vs.cell(row, LAST_COL).value
        if qty is None or price is None or total is None:
            c.ok(False, f"row {row}: К-сть/Ціна/Сума incomplete ({qty!r}, {price!r}, {total!r})")
            continue
        c.ok(
            abs(total - qty * price) < CENT,
            f"row {row} ({label_of(ws, row)[:34]}): Сума {total} != К-сть {qty} × Ціна {price}",
        )

    # Each «Разом …» must equal the rows it claims to sum.
    for row in total_rows:
        formula = str(ws.cell(row, LAST_COL).value or "")
        value = vs.cell(row, LAST_COL).value
        if value is None or not formula.startswith("=SUM(F"):
            continue
        span = formula[len("=SUM(F"):].rstrip(")").split(":F")
        if len(span) != 2 or not all(p.isdigit() for p in span):
            continue
        first, last = int(span[0]), int(span[1])
        expected = sum(
            vs.cell(r, LAST_COL).value or 0
            for r in range(first, last + 1)
        )
        c.ok(
            abs(value - expected) < CENT,
            f"F{row} ({label_of(ws, row)[:34]}): {value} != sum of F{first}:F{last} = {expected}",
        )

    # --- 3. formats, alignment, widths --------------------------------------
    expected_align = {1: "center", 2: "left", 3: "right", 4: "center", 5: "right", 6: "right"}
    for row in line_rows:
        for col, want in expected_align.items():
            got = ws.cell(row, col).alignment.horizontal
            c.ok(got == want, f"{get_column_letter(col)}{row}: alignment {got!r}, expected {want!r}")
        c.ok(
            ws.cell(row, 3).number_format == QTY,
            f"C{row}: К-сть format {ws.cell(row, 3).number_format!r}, expected {QTY!r}",
        )
        for col in (5, LAST_COL):
            fmt = ws.cell(row, col).number_format
            c.ok(
                fmt in (MONEY, MONEY_STRONG),
                f"{get_column_letter(col)}{row}: money format {fmt!r}, expected {MONEY!r}",
            )

    for row in total_rows:
        c.ok(ws.cell(row, LAST_COL).font.bold, f"F{row}: «Разом …» total is not bold")
        border = ws.cell(row, LAST_COL).border
        c.ok(
            border.top.style is not None and border.bottom.style is not None,
            f"F{row}: «Разом …» row has no top/bottom rule",
        )

    # A money column narrower than its widest amount renders as ###.
    for col in (3, 5, LAST_COL):
        letter = get_column_letter(col)
        width = ws.column_dimensions[letter].width or 0
        widest = 0
        for row in line_rows + total_rows:
            value = vs.cell(row, col).value
            if isinstance(value, (int, float)):
                widest = max(widest, len(f"{value:,.2f}".replace(",", " ")))
        c.ok(
            width >= widest,
            f"column {letter}: width {width} < widest value {widest} chars — Excel shows ###",
        )

    print(f"  sheets: {live.sheetnames}")
    print(f"  item rows: {len(line_rows)} | total rows: {len(total_rows)} "
          f"| tables: {len(header_rows)}")
    grand = [r for r in total_rows if label_of(ws, r).startswith("Загальний рахунок")]
    if grand:
        print(f"  Загальний рахунок: {vs.cell(grand[0], LAST_COL).value}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("estimate_id", nargs="?", type=int, default=None)
    parser.add_argument("--out", type=Path, default=Path("/tmp/check_export.xlsx"))
    args = parser.parse_args()

    args.out.parent.mkdir(parents=True, exist_ok=True)
    path, label = build_workbook(args.estimate_id, args.out)
    print(f"Checking {label} -> {path} ({path.stat().st_size:,} bytes)")

    c = Checker()
    check(path, c)
    return c.report()


if __name__ == "__main__":
    raise SystemExit(main())
